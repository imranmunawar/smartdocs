import copy
import os

from django.db import transaction
from smartdocs_app.models import (
    Question, Section, Template, Option
)
from django.shortcuts import get_object_or_404
from django.conf import settings
import secrets
import string
# from ..s3.client import S3_Client
from ..gcs.client import GCS_Client
from django.utils import timezone

from python_docx_replace import docx_get_keys
from docx import Document

from ..enums import messageEnum
from django.contrib import messages

import re

class TemplateDuplicator:
    def __init__(self, template_id, new_template_name, duplicate_questions, file, user):
        self.template_id = template_id
        self.new_template_name = new_template_name
        self.duplicate_questions = duplicate_questions
        self.file = file
        self.section_mapping = {}
        self.question_mapping = {}
        self.option_mapping = {}
        self.user = user

    def generate_random_hash(self, length=20):
        characters = string.ascii_letters + string.digits
        random_hash = ''.join(secrets.choice(characters) for _ in range(length))
        return random_hash


    def duplicate_section(self, old_section_id, new_template):
        section = Section.objects.get(id=old_section_id)
        new_section = copy.copy(section)
        new_section.pk = None
        new_section.template = new_template
        new_section.save()
        return new_section

    def duplicate_question(self, old_question_id, new_section, new_template):
        question = Question.objects.get(id=old_question_id)
        new_question = copy.copy(question)
        new_question.pk = None
        new_question.section = new_section
        new_question.template = new_template
        new_question.save()
        self.question_mapping[old_question_id] = new_question.id
        return new_question


    def _duplicate_options(self, old_question_id, new_question_id, new_section, new_template):

        old_options = Option.objects.filter(question_id=old_question_id)
        for option in old_options:
            new_option = Option.objects.create(
                question_id = new_question_id,
                name = option.name,
                is_active = option.is_active
            ) 
            for old_child_question in option.child_questions.all():
                if old_child_question.id not in self.question_mapping:
                    new_question = self.duplicate_question(old_child_question.id, new_section, new_template)
                    new_question.parent_question_id = new_question_id
                    new_question.save()
                    self.question_mapping[old_child_question.id] = new_question.id
                    new_option.child_questions.add(new_question)
                    self._duplicate_options(old_child_question.id, new_question.id, new_section, new_template)
                else:
                    new_question_id = self.question_mapping[old_child_question.id]
                    new_question = Question.objects.get(id=new_question_id)
                    new_option.child_questions.add(new_question)


    def duplicate(self):
        original_template = get_object_or_404(Template, pk=self.template_id)
        
        with transaction.atomic():
            new_template = copy.copy(original_template)
            new_template.pk = None
            new_template.name = self.new_template_name
            new_template.is_active = False
            new_template.created_at = timezone.now() 
            new_template.last_modified = timezone.now() 
            new_template.user = self.user
            

            placeholders = self.get_placeholders_from_file(self.file)

            new_template.doc_placeholders = placeholders

            random_hash = self.generate_random_hash()
            unique_filename = f'template_{random_hash}.docx'

            # s3 = S3_Client()
            # file_key = s3.upload_template_file_to_s3(self.file, unique_filename)

            gcs = GCS_Client()
            file_key = gcs.upload_template_file_to_gcs(self.file, unique_filename)

            if file_key:
                new_template.template_file_name = unique_filename
                new_template.save()

                if new_template.is_active == True:

                    left_placeholders = self.match_template_placeholders(new_template.id)

                    if left_placeholders:
                        new_template.is_active = False
                        new_template.save()
            
            for section in Section.objects.filter(template=self.template_id):
                print(f"Creating Duplicate for Section: {section.id}")
                if section.parent_section:
                    if section.parent_section.id not in self.section_mapping:
                        new_parent_section = self.duplicate_section(section.parent_section.id, new_template)
                        self.section_mapping[section.parent_section.id] = new_parent_section.id
                        print(f"Created New Parent Section: {new_parent_section.id}")
                    else:
                        new_parent_section = Section.objects.get(id=self.section_mapping[section.parent_section.id])
                if section.id not in self.section_mapping:
                    new_section = self.duplicate_section(section.id, new_template)
                    self.section_mapping[section.id] = new_section.id
                    print(f"Created New Section: {new_section.id}")
                    if section.parent_section:
                        new_section.parent_section = new_parent_section
                        new_section.save()
                print(f"Section duplication done for section: {section.id}")

                if self.duplicate_questions:
                    for question in Question.objects.filter(section=section.id):
                        print(f"Creating Duplicate for Question: {question.id}")
                        if question.parent_question:
                            if question.parent_question.id not in self.question_mapping:
                                print(question.parent_question.id)
                                new_parent_question = self.duplicate_question(question.parent_question.id, new_section, new_template)
                                # self.question_mapping[question.parent_question.id] = new_parent_question.id
                                self._duplicate_options(question.parent_question.id, new_parent_question.id, new_section, new_template)
                                print(f"Created New Parent Question: {new_parent_question.id}")
                            else:
                                new_parent_question = Question.objects.get(id=self.question_mapping[question.parent_question.id])
                        if question.id not in self.question_mapping:
                            new_question = self.duplicate_question(question.id, new_section, new_template)
                            self._duplicate_options(question.id, new_question.id, new_section, new_template)
                            print(f"Created New Question: {new_question.id}")
                            if question.parent_question:
                                new_question.parent_question = new_parent_question
                                new_question.save()
                        print(f"Question duplication done for question: {question.id}")


    def get_placeholders_from_file(self, file):

        if file.name.endswith('.docx'):

            temp_file_path = '/tmp/' + file.name
            with open(temp_file_path, 'wb+') as temp_file:
                for chunk in file.chunks():
                    temp_file.write(chunk)
            
            doc = Document(temp_file_path)
            placeholders = docx_get_keys(doc)
            placeholders_list = list(placeholders)

            file.seek(0)

            return placeholders_list
        return []
    
    def match_template_placeholders(self, template_id):
    
        template = Template.objects.get(id=template_id)
        template_placeholders = template.doc_placeholders or [] #All Document Placeholders

        print('---> TEMPLATE PLACEHOLDERS', template_placeholders)

        question_placeholders = set(
            placeholder[2:-1] for placeholder in 
            Question.objects.filter(template_id=template_id)
            .values_list('placeholder', flat=True)
        ) #All Question Placeholders

        print('---> QUESTION PLACEHOLDERS', question_placeholders)

        simple_unmatched_placeholders = [
            placeholder for placeholder in template_placeholders
            if placeholder not in question_placeholders
        ]

        print('---> SIMPLE UNMATCHED PLACEHOLDERS', simple_unmatched_placeholders)

        multiple_placeholders = self.merge_variables(simple_unmatched_placeholders)
        
        multiple_unmatched_placeholders = [
            placeholder for placeholder in multiple_placeholders
            if placeholder not in question_placeholders
        ]

        print('---> ALL (SIMPLE + MULTIPLE) UNMATCHED PLACEHOLDERS', multiple_unmatched_placeholders)

        return multiple_unmatched_placeholders
    

    def merge_variables(self, variable_list):
        merged_variables = {}
        placeholder_mapping = {}

        for variable in variable_list:
            if re.match(r'.*_\d+$', variable):
                prefix, suffix = variable.rsplit('_', 1)
                
                if prefix not in merged_variables:
                    merged_variables[prefix] = f"{prefix}_N"
                
                placeholder_mapping[variable] = merged_variables[prefix]
            else:
                placeholder_mapping[variable] = variable
        
        merged_list = list(set(placeholder_mapping.values()))
        
        return merged_list
